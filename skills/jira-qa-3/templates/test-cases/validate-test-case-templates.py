import argparse
from pathlib import Path
import sys

from openpyxl import load_workbook


SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))

from importlib.util import module_from_spec, spec_from_file_location


spec = spec_from_file_location("template_generator", SCRIPT_DIR / "generate-test-case-templates.py")
generator = module_from_spec(spec)
spec.loader.exec_module(generator)


def parse_args():
    parser = argparse.ArgumentParser(description="独立校验标准 Excel 和 Xray CSV 测试用例模板")
    parser.add_argument("xlsx", type=Path, help="待校验 Excel 路径")
    parser.add_argument("csv", type=Path, help="待校验 Xray CSV 路径")
    parser.add_argument("requirement_key", help="目标 Jira 需求 Issue Key")
    return parser.parse_args()


def validate_mapping_sheet(xlsx_path, requirement_key):
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    if "Xray导入映射" not in workbook.sheetnames:
        workbook.close()
        raise ValueError("Excel 缺少 Xray导入映射工作表")
    rows = list(workbook["Xray导入映射"].iter_rows(values_only=True))
    workbook.close()
    if not any(row[0] == "标准版本" and str(row[1]) == generator.XRAY_IMPORT_STANDARD_VERSION for row in rows):
        raise ValueError("Xray导入标准版本不一致")
    if not any(row[0] == "测试数据" and row[1] == "Data" for row in rows):
        raise ValueError("Excel映射必须包含测试数据 → Data")
    if not any(row[0] == "Data" and row[1] == "Data" for row in rows):
        raise ValueError("Excel映射必须包含CSV Data → Xray Data")
    if not any(row[0] == "Requirements" and row[1] == 'Link "Tests"' for row in rows):
        raise ValueError('Excel 映射必须包含 Requirements → Link "Tests"')
    if not any(
        row[0] == "模块/分类"
        and row[1] == "Test Repository Path"
        and "需求名称 / XMind业务层级 / 测试分类" in str(row[2])
        for row in rows
    ):
        raise ValueError("Excel映射必须包含需求名称 / XMind业务层级 / 测试分类路径规则")
    if not any(row[0] == "用例编号" and requirement_key in str(row[2]) for row in rows):
        raise ValueError("Excel 映射中的用例编号示例未使用目标 Jira Key")


def main():
    args = parse_args()
    generator.validate_requirement_key(args.requirement_key)
    generator.validate_outputs(args.xlsx.resolve(), args.csv.resolve(), args.requirement_key)
    excel_rows = generator.read_excel_rows(args.xlsx.resolve())
    for row in excel_rows:
        generator.validate_repository_path(row["模块/分类"])
    validate_mapping_sheet(args.xlsx.resolve(), args.requirement_key)
    print(f"STANDARD_VERSION={generator.XRAY_IMPORT_STANDARD_VERSION}")
    print(f"REQUIREMENT_KEY={args.requirement_key}")
    print("VALIDATION=PASSED")


if __name__ == "__main__":
    main()
