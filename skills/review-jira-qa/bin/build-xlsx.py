#!/usr/bin/env python3
"""cases.json -> Excel 工作簿（openpyxl）。

用法：python3 build-xlsx.py <cases.json> <out.xlsx>
"""

import json
import sys
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = [
    "用例编号",
    "模块/分类",
    "用例标题",
    "前置条件",
    "测试数据",
    "测试步骤",
    "预期结果",
    "优先级",
    "用例类型",
    "是否适合自动化",
    "Requirements",
]


def main():
    if len(sys.argv) != 3:
        print("usage: build-xlsx.py <cases.json> <out.xlsx>", file=sys.stderr)
        return 2
    with open(sys.argv[1], encoding="utf-8") as fh:
        data = json.load(fh)
    cases = data.get("cases", [])
    if not cases:
        print("cases.json 没有用例", file=sys.stderr)
        return 2

    wb = Workbook()
    ws = wb.active
    ws.title = "测试用例"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EAF2FF")
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    for c in cases:
        ws.append(
            [
                c.get("id", ""),
                c.get("module", ""),
                c.get("title", ""),
                c.get("precondition", ""),
                c.get("testData", ""),
                c.get("steps", ""),
                c.get("expected", ""),
                c.get("priority", ""),
                c.get("type", ""),
                c.get("automation", ""),
                c.get("requirements", ""),
            ]
        )

    widths = [18, 32, 38, 40, 28, 50, 40, 10, 14, 16, 16]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:K{ws.max_row}"

    stats = wb.create_sheet("用例统计")
    stats.append(["统计项", "值"])
    stats.append(["用例总数", len(cases)])
    for p in ["P0", "P1", "P2", "P3"]:
        stats.append([f"优先级 {p}", sum(1 for c in cases if c.get("priority") == p)])
    for t, n in sorted(Counter(c.get("type", "") for c in cases).items()):
        stats.append([f"类型 {t}", n])
    for cell in stats[1]:
        cell.font = Font(bold=True)

    mapping = wb.create_sheet("Xray导入映射")
    mapping.append(["Excel 列", "Xray 字段", "说明"])
    mapping_rows = [
        ["用例编号", "External ID", "唯一稳定标识；CSV 同名列"],
        ["模块/分类", "Test Repository Path", "需求名称 / XMind业务层级 / 测试分类"],
        ["用例标题", "Summary", ""],
        ["前置条件", "Description", "与用例编号、模块、是否适合自动化合并写入"],
        ["测试数据", "Data", "CSV Data 映射 Xray Data；缺失时留空"],
        ["测试步骤", "Test Steps", "映射 Xray Action"],
        ["预期结果", "Expected Result", ""],
        ["优先级", "Priority", "P0→Highest，P1→High，P2→Medium，P3→Low"],
        ["用例类型", "Labels", "不作为 Xray 独立字段"],
        ["是否适合自动化", "—", "仅本地统计"],
        ["Requirements", 'Link "Tests"', "每行填写目标 Jira 需求 Issue Key"],
    ]
    for row in mapping_rows:
        mapping.append(row)
    for cell in mapping[1]:
        cell.font = Font(bold=True)
    for i, width in enumerate([22, 24, 60], start=1):
        mapping.column_dimensions[get_column_letter(i)].width = width

    wb.save(sys.argv[2])
    print(f"wrote {sys.argv[2]} rows={len(cases)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
