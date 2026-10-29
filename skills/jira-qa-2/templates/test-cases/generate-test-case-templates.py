import argparse
from collections import Counter
import csv
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


XRAY_IMPORT_STANDARD_VERSION = "3"
EXAMPLE_REQUIREMENT_KEY = "DEMO-123"
EXCEL_HEADERS = [
    "用例编号",
    "模块/分类",
    "用例标题",
    "前置条件",
    "测试数据",
    "测试步骤",
    "预期结果",
    "优先级",
    "用例类型",
    "是否适合自动化",
    "Requirements",
]
CSV_HEADERS = [
    "Summary",
    "Description",
    "Test Repository Path",
    "Priority",
    "Labels",
    "Test Steps",
    "Data",
    "Expected Result",
    "External ID",
    "Requirements",
]
PRIORITY_MAP = {
    "P0": "Highest",
    "P1": "High",
    "P2": "Medium",
    "P3": "Low",
}
STANDARD_CASE_TYPES = [
    "功能测试",
    "接口测试",
    "UI测试",
    "数据校验",
    "权限/安全测试",
    "异常测试",
    "边界值测试",
    "兼容性测试",
    "性能测试",
    "回归测试",
]
ABSTRACT_TITLE_PREFIXES = {"条件组合", "一致性", "UI兼容", "回归", "异常"}
GENERIC_REPOSITORY_PATH_SEGMENTS = {"业务管理", "业务配置", "业务管理 > 业务配置"}
CASE_TEMPLATES = [
    {
        "module": "示例需求 / 订单管理 / 创建订单 / 功能验证",
        "title": "【订单管理】填写合法数据成功创建订单",
        "precondition": "1. 已准备具备订单创建权限且会话有效的账号。\n2. 系统中存在可用的客户和商品数据。",
        "test_data": "客户：示例客户A\n商品：示例商品A\n数量：1",
        "steps": "1. 进入【业务系统 > 订单管理 > 创建订单】页面。\n2. 选择客户和商品，填写合法的必填信息并提交。\n3. 查看订单创建结果。",
        "expected": "1. 页面加载和可操作区域正常展示。\n2. 所填数据通过校验并提交成功。\n3. 订单创建成功并生成唯一订单编号。",
        "priority": "P1",
        "type": "功能测试",
        "automation": "是",
        "labels": "订单管理/核心流程",
    },
    {
        "module": "示例需求 / 订单管理 / 创建订单 / 异常校验",
        "title": "【订单管理】必填字段为空时禁止创建订单",
        "precondition": "1. 已准备具备订单创建权限且会话有效的账号。\n2. 系统中存在可用的商品数据。",
        "test_data": "客户：为空\n商品：示例商品A\n数量：1",
        "steps": "1. 进入【业务系统 > 订单管理 > 创建订单】页面。\n2. 保持客户字段为空，填写其他必填信息并提交。\n3. 查看字段提示和订单结果。",
        "expected": "1. 页面加载和可操作区域正常展示。\n2. 客户字段保持为空，其他字段可填写且提交动作被处理。\n3. 系统提示客户为必填项，订单未创建。",
        "priority": "P2",
        "type": "异常测试",
        "automation": "是",
        "labels": "订单管理/字段校验",
    },
    {
        "module": "示例需求 / 订单管理 / 订单列表 / 权限与安全",
        "title": "【订单管理】无查看权限的账号不能访问订单列表",
        "precondition": "1. 已准备不具备订单查看权限且会话有效的测试账号。\n2. 系统中已存在订单数据。",
        "test_data": "",
        "steps": "1. 使用目标账号检查【业务系统 > 订单管理】入口可见性。\n2. 尝试通过已确认的订单列表页面地址直接访问。\n3. 查看访问结果和数据返回。",
        "expected": "1. 菜单中不展示订单列表入口。\n2. 直接访问被权限控制拦截。\n3. 页面或接口返回明确无权限结果，且不泄露订单数据。",
        "priority": "P2",
        "type": "权限/安全测试",
        "automation": "是",
        "labels": "订单管理/权限控制",
    },
]


def parse_args():
    parser = argparse.ArgumentParser(description="生成标准 Excel 和 Xray CSV 测试用例模板")
    parser.add_argument("--requirement-key", help="真实 Jira 需求 Issue Key，例如 ZY20260061-30")
    parser.add_argument("--requirement-name", help="需求名称，用于正式交付文件名")
    parser.add_argument("--output-dir", type=Path, help="输出目录；正式模式应传入 outputs/<scope>/<JiraKey>/test-cases")
    parser.add_argument("--example", action="store_true", help="生成使用 DEMO-123 的仓库脱敏示例模板")
    return parser.parse_args()


def resolve_config(args):
    output_dir = (args.output_dir or Path(__file__).resolve().parent).resolve()
    if args.example:
        if args.requirement_key or args.requirement_name:
            raise ValueError("示例模式不能同时传入需求编号或需求名称")
        return (
            EXAMPLE_REQUIREMENT_KEY,
            output_dir / "test-case-template.xlsx",
            output_dir / "xray-test-case-import-template.csv",
        )
    if not args.requirement_key or not args.requirement_name:
        raise ValueError("正式模式必须同时提供 --requirement-key 和 --requirement-name")
    if args.requirement_key == EXAMPLE_REQUIREMENT_KEY:
        raise ValueError("正式模式不能使用示例需求编号 DEMO-123")
    validate_requirement_key(args.requirement_key)
    safe_name = sanitize_requirement_name(args.requirement_name)
    if args.output_dir is None:
        raise ValueError("正式模式必须通过 --output-dir 指定 outputs/<scope>/<JiraKey>/test-cases")
    if output_dir.name != "test-cases" or output_dir.parent.name != args.requirement_key:
        raise ValueError("正式输出目录必须为 outputs/<scope>/<JiraKey>/test-cases")
    return (
        args.requirement_key,
        output_dir / f"{args.requirement_key}-{safe_name}-测试用例.xlsx",
        output_dir / f"{args.requirement_key}-{safe_name}-Jira导入.csv",
    )


def validate_requirement_key(requirement_key):
    if not re.fullmatch(r"[A-Z][A-Z0-9_]*-\d+", requirement_key):
        raise ValueError(f"Jira 需求编号格式错误：{requirement_key}")


def sanitize_requirement_name(requirement_name):
    name = requirement_name.strip()
    if not name:
        raise ValueError("需求名称不能为空")
    if re.search(r"[<>:\"/\\|?*]", name):
        raise ValueError(f"需求名称包含文件名非法字符：{name}")
    return name


def numbered_items(text):
    return re.findall(r"(?m)^\s*(\d+)[.、)]\s*.+$", text)


def build_cases(requirement_key):
    cases = []
    for index, template in enumerate(CASE_TEMPLATES, 1):
        case = dict(template)
        case["id"] = f"{requirement_key}-TC{index:03d}"
        case["requirements"] = requirement_key
        case["requirement_name"] = "示例需求"
        validate_case(case, requirement_key)
        cases.append(case)
    if len({case["id"] for case in cases}) != len(cases):
        raise ValueError("用例编号必须唯一")
    return cases


def repository_path_parts(path):
    return [part.strip() for part in path.split("/") if part.strip()]


def validate_repository_path(path, requirement_name=None):
    parts = repository_path_parts(path)
    if len(parts) < 3:
        raise ValueError(f"模块/分类必须包含需求名称、XMind业务层级和测试分类：{path}")
    if requirement_name and parts[0] != requirement_name:
        raise ValueError(f"模块/分类必须以需求名称作为根路径：{path}")
    if any(part in GENERIC_REPOSITORY_PATH_SEGMENTS or ">" in part for part in parts):
        raise ValueError(f"模块/分类不得包含通用菜单包装层级或 >：{path}")


def validate_case(case, requirement_key):
    expected_id_pattern = rf"{re.escape(requirement_key)}-TC\d{{3}}"
    if not re.fullmatch(expected_id_pattern, case["id"]):
        raise ValueError(f"用例编号必须以需求编号开头：{case['id']}")
    if case["requirements"] != requirement_key:
        raise ValueError(f"Requirements 必须等于需求编号：{case['id']}")
    validate_repository_path(case["module"], case.get("requirement_name"))
    title_match = re.fullmatch(r"【([^】]+)】(.+)", case["title"])
    if not title_match:
        raise ValueError(f"用例标题必须符合【模块名称】具体测试场景：{case['title']}")
    title_prefix, title_scene = title_match.groups()
    if title_prefix in ABSTRACT_TITLE_PREFIXES or not title_scene.strip():
        raise ValueError(f"用例标题必须使用实际模块名称并描述具体测试场景：{case['title']}")
    preconditions = case["precondition"].splitlines()
    if len(preconditions) > 1 and numbered_items(case["precondition"]) != [str(i) for i in range(1, len(preconditions) + 1)]:
        raise ValueError(f"多条前置条件必须连续编号：{case['id']}")
    steps = numbered_items(case["steps"])
    expected = numbered_items(case["expected"])
    sequence = [str(i) for i in range(1, len(steps) + 1)]
    if steps != sequence or expected != sequence:
        raise ValueError(f"测试步骤与预期结果必须连续编号且一一对应：{case['id']}")
    if case["priority"] not in PRIORITY_MAP:
        raise ValueError(f"用例优先级错误：{case['id']}")
    if case["type"] not in STANDARD_CASE_TYPES:
        raise ValueError(f"用例类型不在标准枚举中：{case['id']}")
    if len(steps) > 4:
        raise ValueError(f"常规模板用例不得超过 4 个主要步骤：{case['id']}")
    test_data = case.get("test_data")
    if test_data is None:
        test_data = ""
    if not isinstance(test_data, str):
        raise ValueError(f"测试数据必须为文本或空值：{case['id']}")
    case["test_data"] = test_data.replace("\r\n", "\n").strip()
    case["jira_priority"] = PRIORITY_MAP[case["priority"]]


def style_header(cells, fill, font, border):
    for cell in cells:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border


def write_excel(path, cases, requirement_key):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="B7C9D6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "测试用例"
    worksheet.append(EXCEL_HEADERS)
    for case in cases:
        worksheet.append([
            case["id"],
            case["module"],
            case["title"],
            case["precondition"],
            case["test_data"],
            case["steps"],
            case["expected"],
            case["priority"],
            case["type"],
            case["automation"],
            case["requirements"],
        ])
    style_header(worksheet[1], header_fill, header_font, border)
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
        for index in [0, 7, 8, 9, 10]:
            row[index].alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:K{worksheet.max_row}"
    widths = [24, 38, 48, 42, 30, 54, 54, 10, 18, 16, 22]
    for index, width in enumerate(widths, 1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    worksheet.row_dimensions[1].height = 26
    for index in range(2, worksheet.max_row + 1):
        worksheet.row_dimensions[index].height = 100
    worksheet.sheet_view.showGridLines = False
    case_table = Table(displayName="TemplateTestCaseTable", ref=f"A1:K{worksheet.max_row}")
    case_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(case_table)

    statistics = workbook.create_sheet("用例统计")
    statistics.append(["统计项", "数量", "占比", "建议/说明"])
    statistics.append(["测试用例总数", len(cases), "100.00%", "模板仅为格式示例，不以凑比例为目标"])
    priority_counts = Counter(case["priority"] for case in cases)
    priority_guidance = {
        "P0": "核心业务阻断、严重数据或安全风险",
        "P1": "主要业务路径的代表性核心用例",
        "P2": "重要功能、异常、权限、边界、UI 或兼容候选",
        "P3": "低价值展示、轻微兼容或补充检查",
    }
    for priority in ["P0", "P1", "P2", "P3"]:
        count = priority_counts[priority]
        statistics.append([priority, count, f"{count / len(cases):.2%}", priority_guidance[priority]])
    type_counts = Counter(case["type"] for case in cases)
    for case_type in STANDARD_CASE_TYPES:
        count = type_counts[case_type]
        statistics.append([case_type, count, f"{count / len(cases):.2%}", "标准中文主要类型；未使用时显示 0"])
    statistics.append(["类型统计合计", sum(type_counts.values()), f"{sum(type_counts.values()) / len(cases):.2%}", "必须等于测试用例总数"])
    statistics.append(["适合自动化", sum(case["automation"] == "是" for case in cases), "", ""])
    statistics.append(["不适合自动化", sum(case["automation"] == "否" for case in cases), "", ""])
    style_header(statistics[1], header_fill, header_font, border)
    for row in statistics.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
    statistics.column_dimensions["A"].width = 28
    statistics.column_dimensions["B"].width = 14
    statistics.column_dimensions["C"].width = 14
    statistics.column_dimensions["D"].width = 42
    statistics.freeze_panes = "A2"
    statistics.sheet_view.showGridLines = False

    mapping = workbook.create_sheet("Xray导入映射")
    mapping_rows = [
        ["配置/CSV字段", "Xray 设置或映射", "说明"],
        ["标准版本", XRAY_IMPORT_STANDARD_VERSION, "Xray导入规范版本：固定10列、Data独立映射、需求级Repository Path"],
        ["用例编号", "External ID", f"必须以需求编号开头：{requirement_key}-TC001"],
        ["Requirements", 'Link "Tests"', f"每条用例固定填写 {requirement_key}"],
        ["需求关联方向", "tests / tested by", "Test 侧显示 tests，需求侧显示 tested by；不得使用 relates to"],
        ["模块/分类", "Test Repository Path", "可变层级：需求名称 / XMind业务层级 / 测试分类；过滤业务管理、业务配置等通用菜单包装层级"],
        ["用例标题", "Summary", "格式为【具体业务模块】具体测试场景"],
        ["前置条件", "Description", "多条前置条件使用连续序号分行；文本写入 Description"],
        ["测试数据", "Data", "Excel测试数据对应CSV Data并映射Xray Data；未提供时保持空白"],
        ["用例类型", "不映射", "Excel 使用单个中文主要类型；Xray CSV 不新增 type-* 或其他类型映射"],
        ["优先级", "Priority", "P0/P1/P2/P3 分别映射 Highest/High/Medium/Low"],
        ["CSV 分隔符", ",", "英文逗号"],
        ["List Delimiter", "/", "Labels 单元格内使用 / 分隔多个标签"],
        ["文件编码", "UTF-8 BOM", "避免中文乱码"],
        ["Summary", "概要", "用例标题"],
        ["Description", "描述", "包含编号、模块、自动化建议和前置条件，不承载主要类型或测试数据"],
        ["Test Repository Path", "Test Repository Path", "与Excel模块/分类一致；以需求名称为根，继承XMind业务层级，以测试分类结尾"],
        ["Priority", "优先级", "Highest/High/Medium/Low"],
        ["Labels", "标签", "多个标签使用 / 分隔"],
        ["Test Steps", "Action", "连续编号的手动测试步骤"],
        ["Data", "Data", "来自Excel测试数据及XMind notes可选测试数据段；未提供时为空"],
        ["Expected Result", "Expected Result", "与测试步骤编号和数量逐项对应"],
        ["External ID", "测试用例标识符", "唯一稳定标识，不能替代 Requirements"],
    ]
    for row in mapping_rows:
        mapping.append(row)
    style_header(mapping[1], header_fill, header_font, border)
    for row in mapping.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
    mapping.column_dimensions["A"].width = 26
    mapping.column_dimensions["B"].width = 32
    mapping.column_dimensions["C"].width = 82
    mapping.freeze_panes = "A2"
    mapping.sheet_view.showGridLines = False

    workbook.save(path)


def write_csv(path, cases):
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(CSV_HEADERS)
        for case in cases:
            description = (
                f"用例编号：{case['id']}\n"
                f"模块：{case['module']}\n"
                f"是否适合自动化：{case['automation']}\n\n"
                f"前置条件\n{case['precondition']}"
            )
            writer.writerow([
                case["title"],
                description,
                case["module"],
                case["jira_priority"],
                case["labels"],
                case["steps"],
                case["test_data"],
                case["expected"],
                case["id"],
                case["requirements"],
            ])


def read_excel_rows(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["测试用例"]
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()
    if list(rows[0]) != EXCEL_HEADERS:
        raise ValueError("Excel 必需表头不符合标准")
    return [dict(zip(EXCEL_HEADERS, row)) for row in rows[1:] if any(value is not None for value in row)]


def read_csv_rows(path):
    if not path.read_bytes().startswith(b"\xef\xbb\xbf"):
        raise ValueError("CSV 必须使用 UTF-8 BOM")
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)
        if reader.fieldnames != CSV_HEADERS:
            raise ValueError("CSV 必需表头或顺序不符合标准")
        return list(reader)


def validate_statistics(path, cases):
    workbook = load_workbook(path, read_only=True, data_only=True)
    statistics = workbook["用例统计"]
    rows = list(statistics.iter_rows(min_row=2, values_only=True))
    workbook.close()
    values = {row[0]: row[1] for row in rows if row and row[0]}
    type_counts = Counter(case["type"] for case in cases)
    for case_type in STANDARD_CASE_TYPES:
        if values.get(case_type) != type_counts[case_type]:
            raise ValueError(f"Excel 类型统计错误：{case_type}")
    if values.get("类型统计合计") != len(cases):
        raise ValueError("Excel 类型统计合计必须等于用例总数")


def validate_outputs(xlsx_path, csv_path, requirement_key):
    expected_prefix = f"{requirement_key}-"
    example_paths = xlsx_path.name == "test-case-template.xlsx" and csv_path.name == "xray-test-case-import-template.csv"
    if not example_paths and (not xlsx_path.name.startswith(expected_prefix) or not csv_path.name.startswith(expected_prefix)):
        raise ValueError("Excel 和 CSV 文件名必须以目标 Jira 需求编号开头")
    excel_rows = read_excel_rows(xlsx_path)
    csv_rows = read_csv_rows(csv_path)
    validate_statistics(xlsx_path, [
        {"type": row["用例类型"]}
        for row in excel_rows
    ])
    if len(excel_rows) != len(csv_rows):
        raise ValueError("Excel 与 CSV 用例行数不一致")
    excel_ids = [row["用例编号"] for row in excel_rows]
    csv_ids = [row["External ID"] for row in csv_rows]
    if len(set(excel_ids)) != len(excel_ids) or len(set(csv_ids)) != len(csv_ids):
        raise ValueError("External ID 必须非空且唯一")
    if set(excel_ids) != set(csv_ids):
        raise ValueError("Excel 用例编号与 CSV External ID 集合不一致")
    id_pattern = re.compile(rf"{re.escape(requirement_key)}-TC\d{{3}}")
    if any(not case_id or not id_pattern.fullmatch(case_id) for case_id in excel_ids):
        raise ValueError("所有用例编号必须使用 <JiraKey>-TC<三位序号>")
    if any(row["Requirements"] != requirement_key for row in excel_rows):
        raise ValueError("Excel Requirements 必须全部等于目标 Jira 需求编号")
    if any(row["Requirements"] != requirement_key for row in csv_rows):
        raise ValueError("CSV Requirements 必须全部等于目标 Jira 需求编号")
    expected_ids = [f"{requirement_key}-TC{index:03d}" for index in range(1, len(excel_ids) + 1)]
    if excel_ids != expected_ids or csv_ids != expected_ids:
        raise ValueError("用例编号必须从 TC001 开始连续递增，且 Excel/CSV 顺序一致")
    for excel_row, csv_row in zip(excel_rows, csv_rows):
        excel_data = (excel_row["测试数据"] or "").replace("\r\n", "\n").strip()
        csv_data = (csv_row["Data"] or "").replace("\r\n", "\n").strip()
        if excel_data != csv_data:
            raise ValueError(f"Excel测试数据与CSV Data不一致：{excel_row['用例编号']}")
        if "\n\n测试数据\n" in (csv_row["Description"] or ""):
            raise ValueError(f"CSV Description不得重复承载测试数据：{csv_row['External ID']}")


def main():
    requirement_key, xlsx_path, csv_path = resolve_config(parse_args())
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    cases = build_cases(requirement_key)
    write_excel(xlsx_path, cases, requirement_key)
    write_csv(csv_path, cases)
    validate_outputs(xlsx_path, csv_path, requirement_key)
    print(f"STANDARD_VERSION={XRAY_IMPORT_STANDARD_VERSION}")
    print(f"REQUIREMENT_KEY={requirement_key}")
    print(f"XLSX={xlsx_path}")
    print(f"CSV={csv_path}")
    print(f"CASES={len(cases)}")
    print("VALIDATION=PASSED")


if __name__ == "__main__":
    main()
