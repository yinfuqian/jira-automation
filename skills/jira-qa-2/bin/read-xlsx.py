#!/usr/bin/env python3
"""读取测试用例 Excel，输出 JSON 摘要供校验脚本使用。"""

import json
import sys

from openpyxl import load_workbook


def main():
    if len(sys.argv) != 2:
        print("usage: read-xlsx.py <xlsx>", file=sys.stderr)
        return 2
    wb = load_workbook(sys.argv[1], read_only=True, data_only=True)
    out = {"sheets": wb.sheetnames, "headers": [], "count": 0, "firstId": "", "lastId": ""}
    if "测试用例" in wb.sheetnames:
        ws = wb["测试用例"]
        rows = list(ws.iter_rows(values_only=True))
        if rows:
            out["headers"] = [str(x).strip() if x is not None else "" for x in rows[0]]
            data_rows = [r for r in rows[1:] if any(v not in (None, "") for v in r)]
            out["count"] = len(data_rows)
            if data_rows:
                out["firstId"] = str(data_rows[0][0] or "")
                out["lastId"] = str(data_rows[-1][0] or "")
    print(json.dumps(out, ensure_ascii=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
